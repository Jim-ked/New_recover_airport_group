import io
import unittest
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook

from backend.services.data_codec_service import DataCodecError, decode_records, encode_records


class DataCodecTests(unittest.TestCase):
    def test_four_formats_keep_unicode_and_all_columns(self):
        rows = [{"id": "001", "name": "机场<&", "optional": None},
                {"id": "002", "name": "=literal", "optional": "后续字段"}]
        for fmt in ("json", "csv", "xlsx", "xml"):
            with self.subTest(fmt=fmt):
                encoded = encode_records(rows, fmt=fmt, columns=("id", "name", "optional"))
                self.assertIsInstance(encoded, bytes)
                decoded = decode_records(encoded, fmt=fmt)
                self.assertEqual("001", decoded[0]["id"])
                self.assertEqual("机场<&", decoded[0]["name"])
                self.assertEqual("后续字段", decoded[1]["optional"])
                self.assertEqual("=literal", decoded[1]["name"])
                if fmt == "xlsx":
                    self.assertTrue(encoded.startswith(b"PK"))
                    workbook = load_workbook(io.BytesIO(encoded))
                    self.assertEqual("s", workbook.active["B3"].data_type)
                    workbook.close()

    def test_standard_excel_numeric_boolean_and_blank_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "基础数据"
        sheet.append(["id", "quantity", "enabled"])
        sheet.append(["001", 12.5, True])
        sheet.append([None, None, None])
        sheet.append(["002", 0, False])
        stream = io.BytesIO(); workbook.save(stream); workbook.close()
        decoded = decode_records(stream.getvalue(), fmt="xlsx")
        self.assertEqual([{"id": "001", "quantity": 12.5, "enabled": True},
                          {"id": "002", "quantity": 0, "enabled": False}], decoded)

    def test_excel_shared_strings_are_supported(self):
        # Conventional Excel files may store text in sharedStrings instead of inlineStr.
        raw = encode_records([{"id": "001"}], fmt="xlsx", columns=("id",))
        result = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(result, "w") as target:
            for part in source.infolist():
                data = source.read(part.filename)
                if part.filename == "[Content_Types].xml":
                    root = ET.fromstring(data)
                    ET.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                        PartName="/xl/sharedStrings.xml", ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml")
                    data = ET.tostring(root)
                elif part.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                    for index, cell in enumerate(root.iter(ns + "c")):
                        for child in list(cell): cell.remove(child)
                        cell.set("t", "s"); ET.SubElement(cell, ns + "v").text = str(index)
                    data = ET.tostring(root)
                target.writestr(part.filename, data)
            target.writestr("xl/sharedStrings.xml", '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>id</t></si><si><t>001</t></si></sst>')
        self.assertEqual([{"id": "001"}], decode_records(result.getvalue(), fmt="xlsx"))

    def test_rejects_malformed_empty_and_duplicate_structure(self):
        cases = [(b"", "json"), (b"<unrelated/>", "xml"), (b"<records>", "xml"),
                 (b"<records><record><id>1</id><id>2</id></record></records>", "xml"),
                 (b"<records><record/></records>", "xml"), (b"id,id\n1,2\n", "csv"),
                 (b"id,name\n1,a,extra\n", "csv"), (b"not excel", "xlsx")]
        for data, fmt in cases:
            with self.subTest(fmt=fmt, data=data):
                with self.assertRaises(DataCodecError): decode_records(data, fmt=fmt)

    def test_formula_input_is_rejected(self):
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["id", "quantity"]); sheet.append(["001", "=1+2"])
        stream = io.BytesIO(); workbook.save(stream); workbook.close()
        with self.assertRaisesRegex(DataCodecError, "formula"):
            decode_records(stream.getvalue(), fmt="xlsx")

    def test_csv_utf8_bom_is_supported(self):
        self.assertEqual([{"id": "001", "name": "机场"}],
                         decode_records("\ufeffid,name\n001,机场\n".encode(), fmt="csv"))
