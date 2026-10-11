"""V1.4.1 workbook contracts; all databases are isolated in temporary directories."""
import hashlib
import sqlite3
import tempfile
import unittest
from contextlib import closing, redirect_stdout
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tools.scenario_db_tool.excel_reader import SCHEMAS, read_excel
from tools.scenario_db_tool.validator import validate
from tools.scenario_db_tool.compiler import build, _rename_exclusive
from tools.scenario_db_tool.models import ScenarioToolError
from tools.scenario_db_tool.cli import main
from backend.storage.situation_repository import SituationRepository
from backend.storage.user_repository import UserRepository


class ScenarioToolTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.default = Path(__file__).resolve().parents[2] / 'runtime/db/airport_group.sqlite3'
        cls.before = cls.default_fingerprint()

    @classmethod
    def default_fingerprint(cls):
        return (hashlib.sha256(cls.default.read_bytes()).hexdigest(), cls.default.stat().st_mtime_ns) if cls.default.exists() else None

    @classmethod
    def tearDownClass(cls):
        if cls.before != cls.default_fingerprint():
            raise AssertionError('默认数据库在工具测试前后发生变化')

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.xlsx = self.root / 'scene.xlsx'
        w = Workbook()
        w.remove(w.active)
        for name, schema in SCHEMAS.items():
            s = w.create_sheet(name)
            s.append([name])
            s.append(list(schema))
        s = w.create_sheet('00_场景说明', 0)
        s.append(['场景说明'])
        s.append([])
        s.append(['字段', '值', '说明'])
        s.append(['场景库编号', 'SCN_TEST'])
        s.append(['场景名称', '测试场景'])
        s.append(['时间窗长度（min）', 15])
        w.create_sheet('_字典_勿改')
        rows = {
            '01_机场': [['AP1', '甲机场', '中型机场', '军用', 120, 30, '否', None, None, None, None, None, '有跑道']],
            '01-1_跑道': [['甲机场', '03/21', 2400, 45, '混凝土', '是']],
            '02_机型': [['AC1', '甲机型', 800, 2000, .15, 1, 1], ['AC2', '乙机型', 800, 2000, .15, 1, 1]],
            '02_资源': [['RS1', '燃油', '燃油', 'L'], ['RS2', '弹药', '航弹', '枚']],
            '02-1_机型资源消耗': [['甲机型', '燃油', 100, '每小时']],
            '03_机场保障': [['AP1', '甲机场', 8]],
            '03-1_机场机型保障': [['甲机场', '甲机型', 10, 2], ['甲机场', '乙机型', 5, 2]],
            '03-2_机场资源库存': [['甲机场', '燃油', 1000, 100, 2000], ['甲机场', '弹药', 50, None, None]],
            '03-3_资源补给': [['甲机场', '燃油', 1, 3, 50]],
            '04_任务': [['MS1', '甲任务', '测试区', 121, 31, 1, 8]],
            '04-1_任务机型需求': [['甲任务', '甲机型', 2, 1]],
            '05_损毁方案': [['DS1', '甲方案', '测试']],
            '05-1_损毁事件': [
                ['是', '甲方案', 'AIR1', 1, '甲机场', '航空器损失', '机场', None, 2, 3, '不恢复', None, None, None, '甲机型', 1],
                ['是', '甲方案', 'AIR1', 1, '甲机场', '航空器损失', '机场', None, 2, 3, '不恢复', None, None, None, '乙机型', 2],
                ['是', '甲方案', 'RES1', 2, '甲机场', '资源损毁', '机场', None, 3, 4, '不恢复', None, '燃油', 800, None, None],
                ['是', '甲方案', 'CAP1', 3, '甲机场', '起降能力损毁', '机场', None, 4, 6, '立即恢复', 0, None, None, None, None],
            ],
        }
        for name, items in rows.items():
            for row in items:
                w[name].append(row)
        w.save(self.xlsx)
        w.close()

    def tearDown(self):
        self.temp.cleanup()

    def edit(self, sheet, cell, value):
        w = load_workbook(self.xlsx)
        w[sheet][cell] = value
        w.save(self.xlsx)
        w.close()

    def test_name_mapping_and_merged_damage(self):
        result = validate(read_excel(self.xlsx))
        self.assertTrue(result.ok, result.issues)
        scene = result.scenario
        self.assertEqual(scene.requirements[0].aircraft_type_id, 'AC1')
        self.assertEqual(scene.situation.missions[0].aircraft_requirements[0].aircraft_type_id, 'AC1')
        events = scene.situation.damage_scenarios[0].events
        self.assertEqual(len(events), 3)
        self.assertEqual(dict(events[0].effect.aircraft_loss), {'AC1': 1, 'AC2': 2})
        self.assertEqual(events[1].effect.to_dict()['inventory_semantics'], 'physical_inventory_v1')
        self.assertEqual(scene.situation.airports[0].resource_replenishments[0].quantity, 50)
        self.assertIsNone(scene.situation.airports[0].operational_profile.resource_stocks[1].storage_capacity)

    def test_duplicate_name_and_ambiguous_reference(self):
        self.edit('02_机型', 'B4', '甲机型')
        issues = validate(read_excel(self.xlsx)).issues
        self.assertTrue(any('名称重复' in i.message for i in issues))
        self.assertTrue(any('多个' in i.message and i.sheet == '04-1_任务机型需求' for i in issues))

    def test_unknown_name(self):
        self.edit('03-1_机场机型保障', 'B3', '不存在')
        issues = validate(read_excel(self.xlsx)).issues
        self.assertTrue(any(i.sheet == '03-1_机场机型保障' and i.row == 3 and i.field == '机型名称' for i in issues))

    def test_missing_sheet(self):
        w = load_workbook(self.xlsx)
        del w['02_资源']
        w.save(self.xlsx)
        w.close()
        self.assertTrue(any(i.sheet == '02_资源' for i in validate(read_excel(self.xlsx)).issues))

    def test_invalid_window(self):
        self.edit('04_任务', 'G3', 1)
        self.assertFalse(validate(read_excel(self.xlsx)).ok)

    def test_replenishment_exceeds_capacity(self):
        self.edit('03-3_资源补给', 'E3', 101)
        result = validate(read_excel(self.xlsx))
        self.assertTrue(any(i.sheet == '03-3_资源补给' and i.row == 3 and 'capacity' in i.message for i in result.issues))

    def test_build_roundtrip_and_no_overwrite(self):
        out = self.root / 'scenes'
        first = build(self.xlsx, output_dir=out)
        original = first.read_bytes()
        second = build(self.xlsx, output_dir=out)
        self.assertNotEqual(first, second)
        self.assertEqual(original, first.read_bytes())
        self.assertEqual(UserRepository(first).authenticate('admin', 'hao123456')['role'], 'admin')
        situation = SituationRepository(first).get_situation('SCN_TEST')
        self.assertEqual(len(situation.damage_scenarios[0].events), 3)
        with closing(sqlite3.connect(first)) as conn:
            self.assertEqual(conn.execute('PRAGMA integrity_check').fetchone()[0], 'ok')
            self.assertEqual(conn.execute('PRAGMA foreign_key_check').fetchall(), [])
            self.assertNotEqual(conn.execute('SELECT password_hash FROM users').fetchone()[0], 'hao123456')
        self.assertFalse(list(out.glob('.*')))

    def test_build_failure_leaves_no_database(self):
        out = self.root / 'scenes'
        with patch('tools.scenario_db_tool.compiler.SituationRepository.save_situation', side_effect=RuntimeError('injected')):
            with self.assertRaises(ScenarioToolError):
                build(self.xlsx, output_dir=out)
        self.assertEqual(list(out.iterdir()), [])

    def test_invalid_build_creates_nothing(self):
        self.edit('04_任务', 'G3', 1)
        out = self.root / 'scenes'
        with self.assertRaises(ScenarioToolError):
            build(self.xlsx, output_dir=out)
        self.assertFalse(out.exists())

    def test_default_db_hash_and_mtime_unchanged(self):
        default = Path(__file__).resolve().parents[2] / 'runtime/db/airport_group.sqlite3'
        def fingerprint():
            return (hashlib.sha256(default.read_bytes()).hexdigest(), default.stat().st_mtime_ns) if default.exists() else None
        before = fingerprint()
        real_connect = sqlite3.connect
        def guarded_connect(database, *args, **kwargs):
            self.assertNotEqual(Path(database).resolve(), default.resolve())
            return real_connect(database, *args, **kwargs)
        with patch('sqlite3.connect', side_effect=guarded_connect):
            validate(read_excel(self.xlsx))
            build(self.xlsx, output_dir=self.root / 'scenes')
        self.assertEqual(fingerprint(), before)

    def test_validate_cli_never_opens_sqlite_or_creates_output(self):
        with patch('sqlite3.connect', side_effect=AssertionError('validate opened SQLite')):
            with redirect_stdout(StringIO()) as output:
                self.assertEqual(main(['validate', str(self.xlsx)]), 0)
        self.assertIn('"ok": true', output.getvalue())
        self.assertEqual(set(self.root.iterdir()), {self.xlsx})

    def test_build_cli_uses_scene_directory(self):
        with patch('tools.scenario_db_tool.compiler.SCENES_DIR', self.root / 'scenes'):
            with redirect_stdout(StringIO()):
                self.assertEqual(main(['build', str(self.xlsx)]), 0)
        self.assertTrue((self.root / 'scenes/SCN_TEST.sqlite3').exists())

    def test_multiple_resources_merge_in_one_physical_event(self):
        self.edit('03-2_机场资源库存', 'E4', 100)
        w = load_workbook(self.xlsx)
        w['05-1_损毁事件'].append(['是', '甲方案', 'RES1', 2, '甲机场', '资源损毁', '机场', None, 3, 4, '不恢复', None, '弹药', 25, None, None])
        w.save(self.xlsx)
        w.close()
        db = build(self.xlsx, output_dir=self.root / 'scenes')
        event = SituationRepository(db).get_situation('SCN_TEST').damage_scenarios[0].events[1]
        self.assertEqual(dict(event.effect.physical_remaining_quantity), {'RS1': 800, 'RS2': 25})

    def test_physical_loss_without_capacity_reports_domain_constraint(self):
        self.edit('05-1_损毁事件', 'M5', '弹药')
        self.edit('05-1_损毁事件', 'N5', 25)
        issues = validate(read_excel(self.xlsx)).issues
        self.assertTrue(any(i.row == 5 and 'storage_capacity' in i.message for i in issues))

    def test_reordered_children_and_multiple_runways_roundtrip(self):
        w = load_workbook(self.xlsx)
        w['01-1_跑道'].append(['甲机场', '04/22', 2600, 45, '混凝土', '是'])
        s = w['03-1_机场机型保障']
        first, second = list(s.values)[2:4]
        for i, value in enumerate(second, 1):
            s.cell(3, i, value)
        for i, value in enumerate(first, 1):
            s.cell(4, i, value)
        w.save(self.xlsx)
        w.close()
        build(self.xlsx, output_dir=self.root / 'scenes')

    def test_inconsistent_event_rows_are_rejected(self):
        self.edit('05-1_损毁事件', 'I4', 1)
        issues = validate(read_excel(self.xlsx)).issues
        self.assertTrue(any(i.row == 4 and '必须一致' in i.message for i in issues))

    def test_bad_input_and_verification_failure_leave_no_database(self):
        with self.assertRaises(ScenarioToolError):
            build(self.root / 'missing.xlsx', output_dir=self.root / 'scenes')
        with patch('tools.scenario_db_tool.compiler._verify', side_effect=ValueError('injected verification failure')):
            with self.assertRaises(ScenarioToolError):
                build(self.xlsx, output_dir=self.root / 'scenes')
        self.assertEqual(list((self.root / 'scenes').iterdir()), [])

    def test_wrong_consumption_basis_and_overlapping_replenishment(self):
        self.edit('02-1_机型资源消耗', 'D3', '每架次')
        self.assertFalse(validate(read_excel(self.xlsx)).ok)
        self.edit('02-1_机型资源消耗', 'D3', '每小时')
        w = load_workbook(self.xlsx)
        w['03-3_资源补给'].append(['甲机场', '燃油', 2, 4, 50])
        w.save(self.xlsx)
        w.close()
        self.assertTrue(any(i.row == 4 and '重叠' in i.message for i in validate(read_excel(self.xlsx)).issues))

    def test_filename_cannot_use_windows_alternate_data_streams(self):
        self.edit('00_场景说明', 'B4', 'SCN_TEST:variant')
        path = build(self.xlsx, output_dir=self.root / 'scenes')
        self.assertEqual(path.name, 'SCN_TEST_variant.sqlite3')
        self.assertIsNotNone(SituationRepository(path).get_situation('SCN_TEST:variant'))

    def test_atomic_rename_race_never_overwrites_competing_file(self):
        real_rename = _rename_exclusive
        competed = []
        def race(source, destination):
            if not competed:
                destination.write_bytes(b'other build won')
                competed.append(destination)
            return real_rename(source, destination)
        with patch('tools.scenario_db_tool.compiler._rename_exclusive', side_effect=race):
            path = build(self.xlsx, output_dir=self.root / 'scenes')
        self.assertNotEqual(path, competed[0])
        self.assertEqual(competed[0].read_bytes(), b'other build won')

    def test_publish_failure_removes_temp_and_preserves_existing_database(self):
        out = self.root / 'scenes'
        out.mkdir()
        existing = out / 'SCN_TEST.sqlite3'
        existing.write_bytes(b'keep')
        with patch('tools.scenario_db_tool.compiler._rename_exclusive', side_effect=PermissionError('injected publish failure')):
            with self.assertRaises(ScenarioToolError):
                build(self.xlsx, output_dir=out)
        self.assertEqual(list(out.iterdir()), [existing])
        self.assertEqual(existing.read_bytes(), b'keep')

    def test_formula_and_missing_header_are_reported(self):
        self.edit('03-3_资源补给', 'E3', '=1+2')
        self.assertTrue(any(i.field == '每窗补给量' and '公式' in i.message for i in validate(read_excel(self.xlsx)).issues))
        self.edit('03-3_资源补给', 'E2', '总补给量')
        self.assertTrue(any(i.row == 2 and '缺少模板列' in i.message for i in validate(read_excel(self.xlsx)).issues))

    def test_unknown_resource_and_scenario_names_are_rejected(self):
        self.edit('05-1_损毁事件', 'M5', '未知资源')
        self.assertTrue(any(i.row == 5 and i.field == '资源名称' for i in validate(read_excel(self.xlsx)).issues))
        self.edit('05-1_损毁事件', 'B3', '未知方案')
        self.assertTrue(any(i.row == 3 and i.field == '损毁方案' for i in validate(read_excel(self.xlsx)).issues))

    def test_disabled_event_names_still_resolve(self):
        self.edit('05-1_损毁事件', 'A5', '否')
        result = validate(read_excel(self.xlsx))
        self.assertTrue(result.ok, result.issues)
        self.assertEqual(len(result.scenario.situation.damage_scenarios[0].events), 2)
        self.edit('05-1_损毁事件', 'M5', '未知资源')
        self.assertFalse(validate(read_excel(self.xlsx)).ok)

    def test_damage_overlap_reports_introducing_row(self):
        w = load_workbook(self.xlsx)
        w['05-1_损毁事件'].append(['是', '甲方案', 'RES2', 4, '甲机场', '资源损毁', '机场', None, 3, 4, '不恢复', None, '燃油', 700, None, None])
        w.save(self.xlsx)
        w.close()
        issues = validate(read_excel(self.xlsx)).issues
        self.assertTrue(any(i.row == 7 and 'overlapping dynamic inventory' in i.message for i in issues))

    def test_invalid_scene_id_is_located_on_metadata_cell(self):
        self.edit('00_场景说明', 'B4', '../outside')
        issues = validate(read_excel(self.xlsx)).issues
        self.assertTrue(any(i.sheet == '00_场景说明' and i.row == 4 and i.field == '场景库编号' for i in issues))


if __name__ == '__main__':
    unittest.main()
