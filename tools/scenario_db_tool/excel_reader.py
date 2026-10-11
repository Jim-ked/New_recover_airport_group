"""Read cells and provenance only. No repository or SQLite access belongs here."""
from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import ExcelData, ExcelRow, ValidationIssue


# Ordered V1.4.1 headers (asterisks are decoration). Values are intermediate keys.
SCHEMAS = {
    '01_机场': dict(zip(
        ['机场编号', '机场名称', '设施类型', '机场角色', '经度', '纬度', '定期服务', 'ICAO', 'IATA', '所属区域', '城市/地区', '高程(m)', '跑道数据状态'],
        ['airport_id', 'airport_name', 'facility_type', 'role', 'longitude', 'latitude', 'scheduled_service', 'icao_code', 'iata_code', 'region', 'municipality', 'elevation_m', 'runway_state'])),
    '01-1_跑道': dict(zip(['机场名称', '跑道编号', '长度(m)', '宽度(m)', '道面', '灯光'],
        ['airport_name', 'runway_label', 'length_m', 'width_m', 'surface', 'lighted'])),
    '02_机型': dict(zip(['机型编号', '机型名称', '规划巡航速度(km/h)', '最大单程航程(km)', '安全余油比例', '离场容量占用比例', '进场容量占用比例'],
        ['aircraft_type_id', 'name', 'speed_kmh', 'max_range_km', 'reserve_ratio', 'departure_capacity_occupancy_factor', 'arrival_capacity_occupancy_factor'])),
    '02_资源': dict(zip(['资源编号', '资源名称', '资源类别', '单位'], ['resource_type_id', 'name', 'category', 'unit'])),
    '02-1_机型资源消耗': dict(zip(['机型名称', '资源名称', '消耗量', '消耗口径'], ['aircraft_name', 'resource_name', 'quantity', 'basis'])),
    '03_机场保障': dict(zip(['机场编号', '机场名称', '起降容量（架次/15min）'], ['airport_id', 'airport_name', 'capacity_per_window'])),
    '03-1_机场机型保障': dict(zip(['机场名称', '机型名称', '初始部署量', '整备时间窗'], ['airport_name', 'aircraft_name', 'initial_quantity', 'tau_reset_windows'])),
    '03-2_机场资源库存': dict(zip(['机场名称', '资源名称', '初始库存', '单窗最大补给能力', '最大储存能力'], ['airport_name', 'resource_name', 'initial_quantity', 'replenishment_capacity_per_window', 'storage_capacity'])),
    '03-3_资源补给': dict(zip(['机场名称', '资源名称', '开始窗', '结束窗', '每窗补给量'], ['airport_name', 'resource_name', 'start_slot', 'end_slot', 'quantity'])),
    '04_任务': dict(zip(['任务编号', '任务名称', '区域类别', '经度', '纬度', '开始窗', '结束窗'], ['mission_id', 'name', 'region_category', 'longitude', 'latitude', 'window_start_slot', 'window_end_slot'])),
    '04-1_任务机型需求': dict(zip(['任务名称', '机型名称', '需求出动架次', '任务点作业时间窗'], ['mission_name', 'aircraft_name', 'required_sorties', 'tau_work_windows'])),
    '05_损毁方案': dict(zip(['方案编号', '方案名称', '说明'], ['damage_scenario_id', 'name', 'description'])),
    '05-1_损毁事件': dict(zip(['启用', '损毁方案', '事件编号', '顺序', '机场名称', '损毁类型', '目标类型', '目标名称', '开始窗', '结束窗', '恢复模式', '剩余起降容量', '资源名称', '资源剩余量', '机型名称', '损失架次'],
        ['enabled', 'scenario_name', 'event_id', 'sequence', 'airport_name', 'damage_type', 'target_type', 'target_name', 'start_slot', 'end_slot', 'recovery_mode', 'remaining_capacity_per_window', 'resource_name', 'remaining_quantity', 'aircraft_name', 'aircraft_loss'])),
}
REQUIRED_SHEETS = ('00_场景说明', *SCHEMAS, '_字典_勿改')


def _label(value):
    return value.strip().rstrip('*＊').strip() if isinstance(value, str) else value


def _value(cell):
    value = cell.value
    return value.strip() or None if isinstance(value, str) else value


def read_excel(path: str | Path) -> ExcelData:
    data = ExcelData(Path(path))
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except (OSError, BadZipFile, InvalidFileException, ValueError) as exc:
        data.issues.append(ValidationIssue('<workbook>', 0, 'file', str(exc)))
        return data
    try:
        for name in REQUIRED_SHEETS:
            if name not in workbook.sheetnames:
                data.issues.append(ValidationIssue(name, 0, 'sheet', '缺少必需 Sheet'))
        if '00_场景说明' in workbook.sheetnames:
            for cells in workbook['00_场景说明'].iter_rows(min_row=4):
                key = _value(cells[0])
                if key not in {'场景库编号', '场景名称', '时间窗长度（min）'}:
                    continue  # Ignore formula summaries and credential display cells.
                if key in data.metadata:
                    data.issues.append(ValidationIssue('00_场景说明', cells[0].row, key, '重复场景说明字段'))
                data.metadata[key] = _value(cells[1]) if len(cells) > 1 else None
                data.metadata_rows[key] = cells[0].row
                if len(cells) > 1 and cells[1].data_type == 'f':
                    data.issues.append(ValidationIssue('00_场景说明', cells[1].row, key, '输入字段不允许公式'))
        for name, schema in SCHEMAS.items():
            data.tables[name] = []
            if name not in workbook.sheetnames:
                continue
            sheet = workbook[name]
            headers = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True), ())
            columns = {}
            for index, header in enumerate(headers):
                label = _label(header)
                if label is None:
                    continue
                if label not in schema:
                    data.issues.append(ValidationIssue(name, 2, str(label), '未知表头，避免静默丢弃输入'))
                elif schema[label] in columns:
                    data.issues.append(ValidationIssue(name, 2, str(label), '重复表头'))
                else:
                    columns[schema[label]] = index
            for label, key in schema.items():
                if key not in columns:
                    data.issues.append(ValidationIssue(name, 2, label, '缺少模板列'))
            labels = {key: label for label, key in schema.items()}
            for cells in sheet.iter_rows(min_row=3):
                if not any(_value(c) is not None for c in cells):
                    continue
                values = {key: _value(cells[index]) if index < len(cells) else None for key, index in columns.items()}
                row = ExcelRow(name, cells[0].row, values, labels)
                data.tables[name].append(row)
                for key, index in columns.items():
                    if index < len(cells) and cells[index].data_type == 'f':
                        data.issues.append(ValidationIssue(name, row.row, labels[key], '输入字段不允许公式'))
    finally:
        workbook.close()
    return data
