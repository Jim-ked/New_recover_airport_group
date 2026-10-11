from __future__ import annotations

"""File codecs only: domain field conversion belongs to BaseDataImportService.

Adapted from the reference project's data_codec_service. XLSX uses openpyxl and
all encoders return real file bytes; encryption is applied by the file service.
"""

import csv
import io
import json
import xml.etree.ElementTree as ET
from typing import Any, Mapping, Sequence
from zipfile import BadZipFile, ZipFile

SUPPORTED_FORMATS = frozenset({"json", "csv", "xlsx", "xml"})
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 100
MIMETYPES = {
    "json": "application/json; charset=utf-8",
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xml": "application/xml; charset=utf-8",
}


class DataCodecError(ValueError):
    pass


def normalize_format(fmt: str) -> str:
    if not isinstance(fmt, str) or fmt.strip().lower() not in SUPPORTED_FORMATS:
        raise DataCodecError("format must be json, csv, xlsx or xml")
    return fmt.strip().lower()


def _headers(values: Sequence[Any]) -> list[str]:
    if not values or len(values) > MAX_COLUMNS:
        raise DataCodecError("file must have 1 to 100 field columns")
    if any(not isinstance(v, str) or not v.strip() for v in values):
        raise DataCodecError("field names must be nonblank strings")
    headers = [v.strip() for v in values]
    if len(set(headers)) != len(headers):
        raise DataCodecError("duplicate field names")
    return headers


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list, tuple)):
        raise DataCodecError("nested tabular values must use JSON cell columns")
    return str(value)


def encode_records(records: Sequence[Mapping[str, Any]], *, fmt: str,
                   columns: Sequence[str] = ()) -> bytes:
    fmt = normalize_format(fmt)
    if fmt == "json":
        return json.dumps(list(records), ensure_ascii=False, allow_nan=False, indent=2).encode("utf-8")
    keys = list(columns) or list(dict.fromkeys(k for row in records for k in row))
    keys = _headers(keys)
    if any(set(row) - set(keys) for row in records):
        raise DataCodecError("record contains fields outside the export columns")
    if fmt == "csv":
        output = io.StringIO(newline="")
        writer = csv.DictWriter(output, fieldnames=keys)
        writer.writeheader()
        writer.writerows({k: _text(row.get(k)) for k in keys} for row in records)
        return output.getvalue().encode("utf-8-sig")
    if fmt == "xml":
        root = ET.Element("records")
        for row in records:
            record = ET.SubElement(root, "record")
            for key in keys:
                element = ET.SubElement(record, key)
                if row.get(key) is None:
                    element.set("null", "true")
                else:
                    element.text = _text(row[key])
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "records"
    sheet.append(keys)
    for index, row in enumerate(records, start=2):
        for col, key in enumerate(keys, start=1):
            value = row.get(key)
            cell = sheet.cell(index, col, value)
            # Literal text, including leading '=', must never become a formula.
            if isinstance(value, str):
                if len(value) > 32767:
                    raise DataCodecError(f"{key}: Excel cell text exceeds 32767 characters")
                cell.data_type = "s"
    output = io.BytesIO()
    try:
        workbook.save(output)
    finally:
        workbook.close()
    return output.getvalue()


def _decode_csv(text: str) -> list[dict[str, Any]]:
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    headers = _headers(next(reader, []))
    records = []
    for row in reader:
        if not row or all(not value.strip() for value in row):
            continue
        if len(row) != len(headers):
            raise DataCodecError(f"CSV row {reader.line_num}: field count does not match header")
        if len(records) >= MAX_ROWS:
            raise DataCodecError("file exceeds 100000 records")
        records.append(dict(zip(headers, row)))
    return records


def _decode_xml(text: str) -> list[dict[str, Any]]:
    if "<!DOCTYPE" in text.upper() or "<!ENTITY" in text.upper():
        raise DataCodecError("XML DTD and entities are not supported")
    root = ET.fromstring(text)
    if root.tag != "records" or root.attrib or (root.text or "").strip():
        raise DataCodecError("XML root must be <records>")
    records = []
    for index, record in enumerate(root):
        if index >= MAX_ROWS:
            raise DataCodecError("file exceeds 100000 records")
        if record.tag != "record" or record.attrib or (record.text or "").strip():
            raise DataCodecError(f"XML record {index + 1}: expected <record>")
        headers = _headers([field.tag for field in record])
        row = {}
        for key, field in zip(headers, record):
            if len(field) or set(field.attrib) - {"null"} or (field.tail or "").strip():
                raise DataCodecError(f"XML field {key}: expected a scalar value")
            is_null = field.attrib.get("null")
            if is_null not in (None, "true") or (is_null and (field.text or "").strip()):
                raise DataCodecError(f"XML field {key}: invalid null value")
            row[key] = None if is_null else (field.text or "")
        records.append(row)
    return records


def _decode_xlsx(data: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook
    # Bound decompression before openpyxl reads the OOXML parts.
    with ZipFile(io.BytesIO(data)) as archive:
        if sum(part.file_size for part in archive.infolist()) > 100 * 1024 * 1024:
            raise DataCodecError("Excel decompressed size exceeds 100 MB")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) != 1:
            raise DataCodecError("Excel must contain exactly one data worksheet")
        sheet = workbook.worksheets[0]
        if sheet.max_column and sheet.max_column > MAX_COLUMNS:
            raise DataCodecError("Excel exceeds 100 columns")
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
            raise DataCodecError("Excel exceeds 100000 records")
        rows = sheet.iter_rows()
        first = next(rows, ())
        headers = _headers([cell.value for cell in first])
        records = []
        for cells in rows:
            if any(cell.data_type in {"f", "e"} for cell in cells):
                raise DataCodecError(f"Excel row {cells[0].row}: formula/error cells are not supported")
            values = [cell.value for cell in cells]
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
                continue
            if len(records) >= MAX_ROWS or len(values) != len(headers):
                raise DataCodecError("Excel row/column count exceeds the supported structure")
            records.append(dict(zip(headers, values)))
        return records
    finally:
        workbook.close()


def decode_records(data: bytes, *, fmt: str) -> Any:
    fmt = normalize_format(fmt)
    if not isinstance(data, bytes) or not data or len(data) > MAX_FILE_BYTES:
        raise DataCodecError("file is empty or exceeds 50 MB")
    try:
        if fmt == "xlsx":
            return _decode_xlsx(data)
        text = data.decode("utf-8-sig")
        if not text.strip():
            raise DataCodecError("file is empty")
        if fmt == "json":
            return json.loads(text)
        if fmt == "csv":
            return _decode_csv(text)
        return _decode_xml(text)
    except DataCodecError:
        raise
    except (ValueError, UnicodeError, csv.Error, ET.ParseError, BadZipFile, KeyError, OSError) as exc:
        raise DataCodecError(f"invalid {fmt} file: {exc}") from exc
